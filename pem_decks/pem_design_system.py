"""
PEM Teaching Library — Master Design System
============================================
Premium PowerPoint design engine for 30 Pediatric Emergency Medicine decks.
Unified visual identity: elegant, modern, fellowship-grade.
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.chart import XL_CHART_TYPE
import copy

# ── COLOR PALETTE ──────────────────────────────────────────────────────────
# Deep navy + teal accent + warm highlight — medical-academic premium feel
NAVY       = RGBColor(0x0B, 0x1D, 0x3A)   # Primary dark
DARK_NAVY  = RGBColor(0x07, 0x12, 0x26)   # Deeper variant
TEAL       = RGBColor(0x00, 0x7B, 0x83)   # Primary accent
LIGHT_TEAL = RGBColor(0x00, 0xA3, 0xAD)   # Lighter accent
CORAL      = RGBColor(0xE8, 0x4D, 0x4D)   # Alert / red flag
AMBER      = RGBColor(0xF0, 0xA5, 0x00)   # Warning / caution
WHITE      = RGBColor(0xFF, 0xFF, 0xFF)
OFF_WHITE  = RGBColor(0xF5, 0xF6, 0xFA)
LIGHT_GRAY = RGBColor(0xE0, 0xE4, 0xEB)
MED_GRAY   = RGBColor(0x8A, 0x92, 0xA3)
DARK_GRAY  = RGBColor(0x3D, 0x44, 0x55)
BLACK      = RGBColor(0x1A, 0x1A, 0x2E)
GREEN      = RGBColor(0x2E, 0x9E, 0x6B)   # Positive / safe
SOFT_BLUE  = RGBColor(0x4A, 0x90, 0xD9)   # Info accent
PALE_TEAL  = RGBColor(0xE6, 0xF5, 0xF6)   # Light background tint

# ── FONTS ──────────────────────────────────────────────────────────────────
FONT_TITLE    = 'Calibri'
FONT_BODY     = 'Calibri'
FONT_ACCENT   = 'Calibri'

# ── SLIDE DIMENSIONS ──────────────────────────────────────────────────────
SLIDE_WIDTH  = Inches(13.333)
SLIDE_HEIGHT = Inches(7.5)


def create_presentation():
    """Create a new presentation with our standard dimensions."""
    prs = Presentation()
    prs.slide_width = SLIDE_WIDTH
    prs.slide_height = SLIDE_HEIGHT
    return prs


def add_shape(slide, left, top, width, height, fill_color=None, border_color=None, border_width=None):
    """Add a rectangle shape with optional fill and border."""
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left, top, width, height)
    shape.line.fill.background()
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = border_width or Pt(1)
    else:
        shape.line.fill.background()
    return shape


def add_rounded_rect(slide, left, top, width, height, fill_color=None, border_color=None):
    """Add a rounded rectangle."""
    shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
    if fill_color:
        shape.fill.solid()
        shape.fill.fore_color.rgb = fill_color
    else:
        shape.fill.background()
    if border_color:
        shape.line.color.rgb = border_color
        shape.line.width = Pt(1.5)
    else:
        shape.line.fill.background()
    return shape


def set_text(shape, text, font_size=18, color=BLACK, bold=False, italic=False,
             alignment=PP_ALIGN.LEFT, font_name=FONT_BODY, line_spacing=1.2):
    """Set text on a shape with formatting."""
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font_name
    p.alignment = alignment
    p.space_after = Pt(0)
    p.space_before = Pt(0)
    if line_spacing != 1.0:
        p.line_spacing = Pt(font_size * line_spacing)
    return tf


def add_paragraph(text_frame, text, font_size=16, color=BLACK, bold=False,
                  italic=False, alignment=PP_ALIGN.LEFT, font_name=FONT_BODY,
                  space_before=0, space_after=4, bullet=False, level=0):
    """Add a paragraph to an existing text frame."""
    p = text_frame.add_paragraph()
    p.text = text
    p.font.size = Pt(font_size)
    p.font.color.rgb = color
    p.font.bold = bold
    p.font.italic = italic
    p.font.name = font_name
    p.alignment = alignment
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.level = level
    return p


def add_footer_bar(slide, deck_number, deck_short_title, slide_num=None):
    """Add consistent footer bar to slide."""
    # Footer background bar
    bar = add_shape(slide, Inches(0), Inches(7.05), SLIDE_WIDTH, Inches(0.45), fill_color=NAVY)
    # Left text: series branding
    left_tf = slide.shapes.add_textbox(Inches(0.4), Inches(7.08), Inches(5), Inches(0.35))
    tf = left_tf.text_frame
    tf.word_wrap = False
    p = tf.paragraphs[0]
    p.text = f"PEM Fellow Teaching Library  |  {deck_short_title}"
    p.font.size = Pt(9)
    p.font.color.rgb = RGBColor(0x8A, 0xB4, 0xC8)
    p.font.name = FONT_BODY
    p.alignment = PP_ALIGN.LEFT
    # Right text: slide number
    if slide_num:
        right_tf = slide.shapes.add_textbox(Inches(11.5), Inches(7.08), Inches(1.5), Inches(0.35))
        tf2 = right_tf.text_frame
        tf2.word_wrap = False
        p2 = tf2.paragraphs[0]
        p2.text = str(slide_num)
        p2.font.size = Pt(9)
        p2.font.color.rgb = RGBColor(0x8A, 0xB4, 0xC8)
        p2.font.name = FONT_BODY
        p2.alignment = PP_ALIGN.RIGHT


def add_title_slide(prs, deck_number, title, subtitle, deck_short_title):
    """Create premium title slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank layout
    # Full navy background
    bg = add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=NAVY)
    # Teal accent stripe
    add_shape(slide, Inches(0), Inches(3.1), Inches(2.8), Inches(0.06), fill_color=TEAL)
    # Deck number badge
    badge = add_rounded_rect(slide, Inches(0.8), Inches(1.4), Inches(1.6), Inches(0.55), fill_color=TEAL)
    set_text(badge, f"DECK {deck_number:02d} / 30", font_size=13, color=WHITE,
             bold=True, alignment=PP_ALIGN.CENTER, font_name=FONT_ACCENT)
    # Series label
    series_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.15), Inches(10), Inches(0.5))
    tf = series_box.text_frame
    p = tf.paragraphs[0]
    p.text = "PEDIATRIC EMERGENCY MEDICINE  ·  FELLOW TEACHING LIBRARY"
    p.font.size = Pt(11)
    p.font.color.rgb = LIGHT_TEAL
    p.font.name = FONT_ACCENT
    p.font.bold = True
    p.alignment = PP_ALIGN.LEFT
    # Main title
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(3.5), Inches(11), Inches(2.0))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(38)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.alignment = PP_ALIGN.LEFT
    # Subtitle
    sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.4), Inches(10), Inches(0.8))
    tf = sub_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = subtitle
    p.font.size = Pt(16)
    p.font.color.rgb = MED_GRAY
    p.font.name = FONT_BODY
    p.alignment = PP_ALIGN.LEFT
    # Bottom line
    add_shape(slide, Inches(0.8), Inches(6.7), Inches(11.5), Inches(0.02), fill_color=TEAL)
    # Footer info
    foot = slide.shapes.add_textbox(Inches(0.8), Inches(6.85), Inches(11), Inches(0.4))
    tf = foot.text_frame
    p = tf.paragraphs[0]
    p.text = "PEM Fellow Teaching Library  ·  Fellowship-Grade Education"
    p.font.size = Pt(10)
    p.font.color.rgb = MED_GRAY
    p.font.name = FONT_BODY
    return slide


def add_section_divider(prs, section_title, section_subtitle="", deck_number=1, deck_short_title=""):
    """Create a section divider slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Left accent panel
    add_shape(slide, Inches(0), Inches(0), Inches(4.5), SLIDE_HEIGHT, fill_color=NAVY)
    # Teal accent bar
    add_shape(slide, Inches(4.5), Inches(0), Inches(0.08), SLIDE_HEIGHT, fill_color=TEAL)
    # Section title on left
    title_box = slide.shapes.add_textbox(Inches(0.6), Inches(2.5), Inches(3.5), Inches(2.5))
    tf = title_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = section_title
    p.font.size = Pt(30)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.alignment = PP_ALIGN.LEFT
    # Subtitle on right
    if section_subtitle:
        sub_box = slide.shapes.add_textbox(Inches(5.2), Inches(3.0), Inches(7.5), Inches(2.0))
        tf = sub_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = section_subtitle
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
        p.alignment = PP_ALIGN.LEFT
    add_footer_bar(slide, deck_number, deck_short_title)
    return slide


def add_objectives_slide(prs, objectives, deck_number=1, deck_short_title="", slide_num=2):
    """Create learning objectives slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    # Light background
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=OFF_WHITE)
    # Title bar
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.1), fill_color=NAVY)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(10), Inches(0.7))
    set_text(title_box, "Learning Objectives", font_size=28, color=WHITE,
             bold=True, font_name=FONT_TITLE)
    # Objectives
    for i, obj in enumerate(objectives):
        y = 1.5 + i * 0.85
        # Number circle
        circle = slide.shapes.add_shape(MSO_SHAPE.OVAL, Inches(0.8), Inches(y), Inches(0.45), Inches(0.45))
        circle.fill.solid()
        circle.fill.fore_color.rgb = TEAL
        circle.line.fill.background()
        set_text(circle, str(i + 1), font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
        # Text
        obj_box = slide.shapes.add_textbox(Inches(1.5), Inches(y - 0.02), Inches(10.5), Inches(0.6))
        tf = obj_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = obj
        p.font.size = Pt(16)
        p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    return slide


def add_content_slide(prs, title, bullets, notes="", deck_number=1,
                      deck_short_title="", slide_num=3, accent_color=TEAL):
    """Create standard content slide with title and bullet points."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=WHITE)
    # Title area
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.1), fill_color=NAVY)
    # Accent bar under title
    add_shape(slide, Inches(0.8), Inches(1.1), Inches(2.5), Inches(0.05), fill_color=accent_color)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.5), Inches(0.7))
    set_text(title_box, title, font_size=26, color=WHITE, bold=True, font_name=FONT_TITLE)
    # Bullet content
    content_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11.5), Inches(5.2))
    tf = content_box.text_frame
    tf.word_wrap = True
    for i, bullet in enumerate(bullets):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        # Support sub-bullets with ">" prefix
        if bullet.startswith(">"):
            p.text = bullet[1:].strip()
            p.level = 1
            p.font.size = Pt(14)
            p.font.color.rgb = MED_GRAY
        else:
            p.text = bullet
            p.level = 0
            p.font.size = Pt(16)
            p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
        p.space_before = Pt(4)
        p.space_after = Pt(6)
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    # Notes
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_two_column_slide(prs, title, left_title, left_items, right_title, right_items,
                         notes="", deck_number=1, deck_short_title="", slide_num=3,
                         left_color=TEAL, right_color=CORAL):
    """Create a two-column comparison slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=WHITE)
    # Title bar
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.1), fill_color=NAVY)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.5), Inches(0.7))
    set_text(title_box, title, font_size=26, color=WHITE, bold=True, font_name=FONT_TITLE)
    # Left column header
    left_hdr = add_rounded_rect(slide, Inches(0.5), Inches(1.4), Inches(5.8), Inches(0.6), fill_color=left_color)
    set_text(left_hdr, left_title, font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    # Left column content
    left_box = slide.shapes.add_textbox(Inches(0.7), Inches(2.2), Inches(5.4), Inches(4.5))
    tf = left_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(left_items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
        p.space_after = Pt(5)
    # Divider
    add_shape(slide, Inches(6.55), Inches(1.4), Inches(0.03), Inches(5.2), fill_color=LIGHT_GRAY)
    # Right column header
    right_hdr = add_rounded_rect(slide, Inches(6.9), Inches(1.4), Inches(5.8), Inches(0.6), fill_color=right_color)
    set_text(right_hdr, right_title, font_size=16, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    # Right column content
    right_box = slide.shapes.add_textbox(Inches(7.1), Inches(2.2), Inches(5.4), Inches(4.5))
    tf = right_box.text_frame
    tf.word_wrap = True
    for i, item in enumerate(right_items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = item
        p.font.size = Pt(14)
        p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
        p.space_after = Pt(5)
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_key_point_slide(prs, title, point_text, sub_text="", notes="",
                        deck_number=1, deck_short_title="", slide_num=3,
                        bg_color=NAVY, accent=TEAL):
    """Create a high-impact key point / pearl slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=bg_color)
    # "CLINICAL PEARL" or "KEY POINT" label
    label_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.5), Inches(11), Inches(0.5))
    tf = label_box.text_frame
    p = tf.paragraphs[0]
    p.text = f"▎ {title.upper()}"
    p.font.size = Pt(14)
    p.font.color.rgb = accent
    p.font.bold = True
    p.font.name = FONT_ACCENT
    # Main point
    point_box = slide.shapes.add_textbox(Inches(0.8), Inches(2.5), Inches(11.5), Inches(2.5))
    tf = point_box.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = point_text
    p.font.size = Pt(28)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.alignment = PP_ALIGN.LEFT
    # Sub text
    if sub_text:
        sub_box = slide.shapes.add_textbox(Inches(0.8), Inches(5.0), Inches(11), Inches(1.5))
        tf = sub_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = sub_text
        p.font.size = Pt(15)
        p.font.color.rgb = MED_GRAY
        p.font.name = FONT_BODY
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_table_slide(prs, title, headers, rows, notes="", deck_number=1,
                    deck_short_title="", slide_num=3):
    """Create a slide with a formatted table."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=WHITE)
    # Title bar
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.1), fill_color=NAVY)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.5), Inches(0.7))
    set_text(title_box, title, font_size=26, color=WHITE, bold=True, font_name=FONT_TITLE)
    # Table
    num_rows = len(rows) + 1  # +1 for header
    num_cols = len(headers)
    tbl_width = Inches(11.5)
    tbl_height = Inches(min(5.5, 0.5 + len(rows) * 0.5))
    table_shape = slide.shapes.add_table(num_rows, num_cols,
                                          Inches(0.9), Inches(1.4),
                                          tbl_width, tbl_height)
    table = table_shape.table
    # Style header
    for i, h in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = h
        for p in cell.text_frame.paragraphs:
            p.font.size = Pt(13)
            p.font.color.rgb = WHITE
            p.font.bold = True
            p.font.name = FONT_BODY
            p.alignment = PP_ALIGN.CENTER
        cell.fill.solid()
        cell.fill.fore_color.rgb = TEAL
    # Style rows
    for r_idx, row in enumerate(rows):
        for c_idx, val in enumerate(row):
            cell = table.cell(r_idx + 1, c_idx)
            cell.text = str(val)
            for p in cell.text_frame.paragraphs:
                p.font.size = Pt(12)
                p.font.color.rgb = DARK_GRAY
                p.font.name = FONT_BODY
                p.alignment = PP_ALIGN.LEFT
            cell.fill.solid()
            cell.fill.fore_color.rgb = OFF_WHITE if r_idx % 2 == 0 else WHITE
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_algorithm_slide(prs, title, steps, notes="", deck_number=1,
                        deck_short_title="", slide_num=3):
    """Create an algorithm / flowchart-style slide with sequential boxes."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=WHITE)
    # Title bar
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.1), fill_color=NAVY)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.5), Inches(0.7))
    set_text(title_box, title, font_size=26, color=WHITE, bold=True, font_name=FONT_TITLE)
    # Algorithm steps as connected boxes
    colors = [TEAL, SOFT_BLUE, NAVY, TEAL, SOFT_BLUE, NAVY, TEAL, SOFT_BLUE]
    num_steps = len(steps)
    if num_steps <= 5:
        # Vertical layout
        box_h = min(0.7, 4.5 / num_steps)
        gap = 0.25
        start_y = 1.5
        for i, step in enumerate(steps):
            y = start_y + i * (box_h + gap)
            col = colors[i % len(colors)]
            box = add_rounded_rect(slide, Inches(2.5), Inches(y), Inches(8), Inches(box_h), fill_color=col)
            set_text(box, step, font_size=14, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
            # Arrow between boxes
            if i < num_steps - 1:
                arrow_y = y + box_h
                arr = slide.shapes.add_shape(MSO_SHAPE.DOWN_ARROW, Inches(6.3), Inches(arrow_y),
                                              Inches(0.4), Inches(gap - 0.02))
                arr.fill.solid()
                arr.fill.fore_color.rgb = LIGHT_GRAY
                arr.line.fill.background()
    else:
        # Two-column layout for many steps
        box_h = 0.55
        gap = 0.18
        mid = (num_steps + 1) // 2
        for i, step in enumerate(steps):
            col = colors[i % len(colors)]
            if i < mid:
                x = 0.5
                y = 1.5 + i * (box_h + gap)
            else:
                x = 6.9
                y = 1.5 + (i - mid) * (box_h + gap)
            box = add_rounded_rect(slide, Inches(x), Inches(y), Inches(5.8), Inches(box_h), fill_color=col)
            set_text(box, f"{i+1}. {step}", font_size=12, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_case_slide(prs, case_title, scenario, questions=None, notes="",
                   deck_number=1, deck_short_title="", slide_num=3):
    """Create a clinical case slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=OFF_WHITE)
    # Title bar with case accent
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.1), fill_color=DARK_NAVY)
    # Case label
    badge = add_rounded_rect(slide, Inches(0.8), Inches(0.25), Inches(1.8), Inches(0.5), fill_color=AMBER)
    set_text(badge, "CLINICAL CASE", font_size=12, color=WHITE, bold=True, alignment=PP_ALIGN.CENTER)
    title_box = slide.shapes.add_textbox(Inches(2.9), Inches(0.2), Inches(9.5), Inches(0.7))
    set_text(title_box, case_title, font_size=24, color=WHITE, bold=True, font_name=FONT_TITLE)
    # Scenario box
    scenario_box = add_rounded_rect(slide, Inches(0.6), Inches(1.4), Inches(12), Inches(2.8),
                                     fill_color=WHITE, border_color=LIGHT_GRAY)
    tf = scenario_box.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.3)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.2)
    p = tf.paragraphs[0]
    p.text = scenario
    p.font.size = Pt(15)
    p.font.color.rgb = DARK_GRAY
    p.font.name = FONT_BODY
    # Questions
    if questions:
        q_box = slide.shapes.add_textbox(Inches(0.8), Inches(4.5), Inches(11.5), Inches(2.2))
        tf = q_box.text_frame
        tf.word_wrap = True
        for i, q in enumerate(questions):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = f"▸  {q}"
            p.font.size = Pt(14)
            p.font.color.rgb = TEAL
            p.font.bold = True
            p.font.name = FONT_BODY
            p.space_after = Pt(6)
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_pitfalls_slide(prs, title, pitfalls, notes="", deck_number=1,
                       deck_short_title="", slide_num=3):
    """Create a pitfalls / red flags slide with warning styling."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=WHITE)
    # Title bar - coral/red for warning
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.1), fill_color=CORAL)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.5), Inches(0.7))
    set_text(title_box, title, font_size=26, color=WHITE, bold=True, font_name=FONT_TITLE)
    # Pitfall items
    for i, pitfall in enumerate(pitfalls):
        y = 1.4 + i * 0.78
        # Warning icon area
        icon_box = add_rounded_rect(slide, Inches(0.5), Inches(y), Inches(0.5), Inches(0.5),
                                     fill_color=RGBColor(0xFD, 0xEB, 0xEB))
        set_text(icon_box, "⚠", font_size=16, color=CORAL, alignment=PP_ALIGN.CENTER)
        # Text
        pit_box = slide.shapes.add_textbox(Inches(1.2), Inches(y), Inches(11.5), Inches(0.6))
        tf = pit_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = pitfall
        p.font.size = Pt(15)
        p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    if notes:
        slide.notes_slide.notes_text_frame.text = notes
    return slide


def add_takeaway_slide(prs, takeaways, deck_number=1, deck_short_title="", slide_num=20):
    """Create take-home messages slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=NAVY)
    # Label
    label = slide.shapes.add_textbox(Inches(0.8), Inches(0.8), Inches(11), Inches(0.5))
    tf = label.text_frame
    p = tf.paragraphs[0]
    p.text = "TAKE-HOME MESSAGES"
    p.font.size = Pt(14)
    p.font.color.rgb = TEAL
    p.font.bold = True
    p.font.name = FONT_ACCENT
    # Accent line
    add_shape(slide, Inches(0.8), Inches(1.5), Inches(3), Inches(0.04), fill_color=TEAL)
    # Messages
    for i, msg in enumerate(takeaways):
        y = 1.9 + i * 0.85
        num_box = slide.shapes.add_textbox(Inches(0.8), Inches(y), Inches(0.5), Inches(0.5))
        tf = num_box.text_frame
        p = tf.paragraphs[0]
        p.text = str(i + 1)
        p.font.size = Pt(22)
        p.font.color.rgb = TEAL
        p.font.bold = True
        p.font.name = FONT_ACCENT
        msg_box = slide.shapes.add_textbox(Inches(1.5), Inches(y), Inches(11), Inches(0.7))
        tf = msg_box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = msg
        p.font.size = Pt(17)
        p.font.color.rgb = WHITE
        p.font.name = FONT_BODY
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    return slide


def add_references_slide(prs, references, deck_number=1, deck_short_title="", slide_num=25):
    """Create references slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=WHITE)
    # Title bar
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, Inches(1.1), fill_color=NAVY)
    title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.2), Inches(11.5), Inches(0.7))
    set_text(title_box, "References", font_size=26, color=WHITE, bold=True, font_name=FONT_TITLE)
    # References text
    ref_box = slide.shapes.add_textbox(Inches(0.8), Inches(1.4), Inches(11.5), Inches(5.3))
    tf = ref_box.text_frame
    tf.word_wrap = True
    for i, ref in enumerate(references):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.text = f"{i+1}. {ref}"
        p.font.size = Pt(11)
        p.font.color.rgb = DARK_GRAY
        p.font.name = FONT_BODY
        p.space_after = Pt(4)
    add_footer_bar(slide, deck_number, deck_short_title, slide_num)
    return slide


def add_end_slide(prs, deck_number=1, deck_short_title=""):
    """Create closing slide."""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    add_shape(slide, Inches(0), Inches(0), SLIDE_WIDTH, SLIDE_HEIGHT, fill_color=NAVY)
    # Thank you
    ty_box = slide.shapes.add_textbox(Inches(2), Inches(2.5), Inches(9), Inches(1.5))
    tf = ty_box.text_frame
    p = tf.paragraphs[0]
    p.text = "Questions & Discussion"
    p.font.size = Pt(36)
    p.font.color.rgb = WHITE
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.alignment = PP_ALIGN.CENTER
    # Sub
    sub = slide.shapes.add_textbox(Inches(2), Inches(4.2), Inches(9), Inches(1))
    tf = sub.text_frame
    p = tf.paragraphs[0]
    p.text = "PEM Fellow Teaching Library"
    p.font.size = Pt(16)
    p.font.color.rgb = LIGHT_TEAL
    p.font.name = FONT_BODY
    p.alignment = PP_ALIGN.CENTER
    # Teal line
    add_shape(slide, Inches(5), Inches(4.0), Inches(3.3), Inches(0.04), fill_color=TEAL)
    return slide


def save_deck(prs, filepath):
    """Save the presentation."""
    prs.save(filepath)
    return filepath
